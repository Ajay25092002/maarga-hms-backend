
# Router/PricingImport.py

import logging
import os
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

import psycopg2
from psycopg2.extras import RealDictCursor
from openpyxl import load_workbook

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse


logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/api/pricing/import",
    tags=["Pricing Import"],
)

TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent
    / "Templates"
    / "OP_IP_Pricing_Import_Template.xlsx"
)

REQUIRED_SHEETS = {
    "ServiceGroups": [
        "Service Group Name",
        "Is Active",
    ],
    "Services": [
        "Service Group Name",
        "Service Name",
        "Available for OP",
        "Available for IP",
        "Is Active",
    ],
    "RoomTypes": [
        "Room Type Name",
        "Is Active",
    ],
    "Doctors": [
        "Doctor Name",
        "Is Active",
    ],
    "ServicePrices": [
        "Service Group Name",
        "Service Name",
        "Bill Type",
        "Room Type",
        "Doctor Name",
        "Amount",
        "Is Active",
    ],
}


# ---------------------------------------------------------
# DATABASE CONNECTION
# ---------------------------------------------------------

def get_db_connection():
    database_url = os.getenv("DATABASE_URL")

    if not database_url:
        raise HTTPException(
            status_code=500,
            detail="DATABASE_URL is not configured.",
        )

    try:
        return psycopg2.connect(database_url, options="-c timezone=Asia/Kolkata")

    except psycopg2.Error:
        logger.exception("Database connection failed.")

        raise HTTPException(
            status_code=500,
            detail="Unable to connect to the database.",
        )


# ---------------------------------------------------------
# GENERAL HELPERS
# ---------------------------------------------------------

def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def parse_yes_no(value, default=None):
    value = clean_text(value).lower()

    if not value and default is not None:
        return default

    if value in ("yes", "y", "true", "1"):
        return True

    if value in ("no", "n", "false", "0"):
        return False

    return None


def parse_amount(value):
    if value is None or clean_text(value) == "":
        return None

    try:
        amount = Decimal(str(value).strip())

    except (InvalidOperation, ValueError):
        return None

    if not amount.is_finite() or amount < 0:
        return None

    if amount.as_tuple().exponent < -2:
        return None

    if amount >= Decimal("10000000000"):
        return None

    return amount.quantize(Decimal("0.01"))


def normalized(value):
    return clean_text(value).casefold()


def is_blank_row(values):
    return all(value is None or clean_text(value) == "" for value in values)


# ---------------------------------------------------------
# READ AND VALIDATE EXCEL WORKBOOK
# ---------------------------------------------------------

def read_workbook(file_bytes):
    errors = []
    data = {}

    try:
        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )

    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unable to read the Excel workbook. "
                "Please upload a valid .xlsx template."
            ),
        )

    try:
        missing_sheets = [
            name
            for name in REQUIRED_SHEETS
            if name not in workbook.sheetnames
        ]

        if missing_sheets:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": "Required worksheets are missing.",
                    "missing_sheets": missing_sheets,
                },
            )

        for sheet_name, expected_headers in REQUIRED_SHEETS.items():
            sheet = workbook[sheet_name]

            if sheet.max_row > 10002:
                errors.append({
                    "sheet": sheet_name,
                    "row": None,
                    "message": "Maximum supported rows exceeded.",
                })
                continue

            headers = [
                clean_text(cell)
                for cell in next(
                    sheet.iter_rows(
                        min_row=1,
                        max_row=1,
                        values_only=True,
                    ),
                    (),
                )
            ]

            if headers != expected_headers:
                errors.append({
                    "sheet": sheet_name,
                    "row": 1,
                    "message": (
                        "Column headers do not match the template. "
                        f"Expected: {expected_headers}"
                    ),
                })
                continue

            records = []

            for row_number, values in enumerate(
                sheet.iter_rows(
                    min_row=2,
                    values_only=True,
                ),
                start=2,
            ):
                values = list(values[:len(expected_headers)])

                if is_blank_row(values):
                    continue

                record = dict(zip(expected_headers, values))
                record["_row"] = row_number
                records.append(record)

            data[sheet_name] = records

    finally:
        workbook.close()

    return data, errors


# ---------------------------------------------------------
# WORKBOOK BUSINESS VALIDATION
# ---------------------------------------------------------

def validate_workbook(data):
    errors = []

    groups = {}
    services = {}
    room_types = {}
    doctors = {}
    price_conditions = set()

    # Validate service groups.
    for row in data.get("ServiceGroups", []):
        name = clean_text(row["Service Group Name"])
        active = parse_yes_no(row["Is Active"], default=True)

        if not name:
            errors.append({
                "sheet": "ServiceGroups",
                "row": row["_row"],
                "message": "Service Group Name is required.",
            })
            continue

        if active is None:
            errors.append({
                "sheet": "ServiceGroups",
                "row": row["_row"],
                "message": "Is Active must be Yes or No.",
            })
            continue

        key = normalized(name)

        if key in groups:
            errors.append({
                "sheet": "ServiceGroups",
                "row": row["_row"],
                "message": "Duplicate service group name.",
            })
            continue

        groups[key] = {
            "name": name,
            "is_active": active,
            "row": row["_row"],
        }

    # Validate room types.
    for row in data.get("RoomTypes", []):
        name = clean_text(row["Room Type Name"])
        active = parse_yes_no(row["Is Active"], default=True)

        if not name:
            errors.append({
                "sheet": "RoomTypes",
                "row": row["_row"],
                "message": "Room Type Name is required.",
            })
            continue

        if active is None:
            errors.append({
                "sheet": "RoomTypes",
                "row": row["_row"],
                "message": "Is Active must be Yes or No.",
            })
            continue

        key = normalized(name)

        if key in room_types:
            errors.append({
                "sheet": "RoomTypes",
                "row": row["_row"],
                "message": "Duplicate room type name.",
            })
            continue

        room_types[key] = {
            "name": name,
            "is_active": active,
            "row": row["_row"],
        }

    # Validate doctors.
    for row in data.get("Doctors", []):
        name = clean_text(row["Doctor Name"])
        active = parse_yes_no(row["Is Active"], default=True)

        if not name:
            errors.append({
                "sheet": "Doctors",
                "row": row["_row"],
                "message": "Doctor Name is required.",
            })
            continue

        if active is None:
            errors.append({
                "sheet": "Doctors",
                "row": row["_row"],
                "message": "Is Active must be Yes or No.",
            })
            continue

        key = normalized(name)

        if key in doctors:
            errors.append({
                "sheet": "Doctors",
                "row": row["_row"],
                "message": "Duplicate doctor name.",
            })
            continue

        doctors[key] = {
            "name": name,
            "is_active": active,
            "row": row["_row"],
        }

    # Validate services.
    for row in data.get("Services", []):
        group_name = clean_text(row["Service Group Name"])
        service_name = clean_text(row["Service Name"])
        available_op = parse_yes_no(row["Available for OP"])
        available_ip = parse_yes_no(row["Available for IP"])
        active = parse_yes_no(row["Is Active"], default=True)

        if not group_name or not service_name:
            errors.append({
                "sheet": "Services",
                "row": row["_row"],
                "message": (
                    "Service Group Name and Service Name are required."
                ),
            })
            continue

        if available_op is None or available_ip is None:
            errors.append({
                "sheet": "Services",
                "row": row["_row"],
                "message": (
                    "Available for OP and Available for IP "
                    "must each be Yes or No."
                ),
            })
            continue

        if not available_op and not available_ip:
            errors.append({
                "sheet": "Services",
                "row": row["_row"],
                "message": (
                    "A service must be available for OP, IP, or both."
                ),
            })
            continue

        if active is None:
            errors.append({
                "sheet": "Services",
                "row": row["_row"],
                "message": "Is Active must be Yes or No.",
            })
            continue

        group_key = normalized(group_name)
        service_key = (group_key, normalized(service_name))

        if group_key in groups and not groups[group_key]["is_active"]:
            errors.append({
                "sheet": "Services",
                "row": row["_row"],
                "message": "The selected service group is inactive.",
            })
            continue

        if service_key in services:
            errors.append({
                "sheet": "Services",
                "row": row["_row"],
                "message": (
                    "Duplicate service name within the same group."
                ),
            })
            continue

        services[service_key] = {
            "group_name": group_name,
            "name": service_name,
            "available_for_op": available_op,
            "available_for_ip": available_ip,
            "is_active": active,
            "row": row["_row"],
        }

    # Validate pricing records.
    for row in data.get("ServicePrices", []):
        group_name = clean_text(row["Service Group Name"])
        service_name = clean_text(row["Service Name"])
        bill_type = clean_text(row["Bill Type"]).upper()
        room_name = clean_text(row["Room Type"])
        doctor_name = clean_text(row["Doctor Name"])
        amount = parse_amount(row["Amount"])
        active = parse_yes_no(row["Is Active"], default=True)

        if not group_name or not service_name:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": (
                    "Service Group Name and Service Name are required."
                ),
            })
            continue

        if bill_type not in ("OP", "IP"):
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": "Bill Type must be OP or IP.",
            })
            continue

        if amount is None:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": (
                    "Amount must be a non-negative number "
                    "with no more than two decimal places."
                ),
            })
            continue

        if active is None:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": "Is Active must be Yes or No.",
            })
            continue

        if room_name and doctor_name:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": (
                    "A price cannot be both room-specific "
                    "and doctor-specific."
                ),
            })
            continue

        if bill_type == "OP" and room_name:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": "Room-specific prices are only valid for IP.",
            })
            continue

        if bill_type == "IP" and doctor_name:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": "Doctor-specific prices are only valid for OP.",
            })
            continue

        group_key = normalized(group_name)
        service_key = (group_key, normalized(service_name))

        service = services.get(service_key)

        if service is None:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": (
                    "The service must be listed in the Services sheet "
                    "under the selected service group."
                ),
            })
            continue

        if not service["is_active"] and active:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": "Cannot activate a price for an inactive service.",
            })
            continue

        if bill_type == "OP" and not service["available_for_op"]:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": "This service is not available for OP billing.",
            })
            continue

        if bill_type == "IP" and not service["available_for_ip"]:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": "This service is not available for IP billing.",
            })
            continue

        if room_name:
            room_key = normalized(room_name)

            if room_key not in room_types:
                errors.append({
                    "sheet": "ServicePrices",
                    "row": row["_row"],
                    "message:": (
                        "Room Type must exist in the RoomTypes sheet."
                    ),
                })
                continue

            if not room_types[room_key]["is_active"] and active:
                errors.append({
                    "sheet": "ServicePrices",
                    "row": row["_row"],
                    "message": "Cannot activate a price for an inactive room type.",
                })
                continue

        if doctor_name:
            doctor_key = normalized(doctor_name)

            if doctor_key not in doctors:
                errors.append({
                    "sheet": "ServicePrices",
                    "row": row["_row"],
                    "message": "Doctor Name must exist in the Doctors sheet.",
                })
                continue

            if not doctors[doctor_key]["is_active"] and active:
                errors.append({
                    "sheet": "ServicePrices",
                    "row": row["_row"],
                    "message": "Cannot activate a price for an inactive doctor.",
                })
                continue

        condition = (
            service_key,
            bill_type,
            normalized(room_name),
            normalized(doctor_name),
        )

        if condition in price_conditions:
            errors.append({
                "sheet": "ServicePrices",
                "row": row["_row"],
                "message": (
                    "Duplicate service, bill type, room type, "
                    "and doctor pricing condition."
                ),
            })
            continue

        price_conditions.add(condition)

    return errors


# ---------------------------------------------------------
# DATABASE UPSERT HELPERS
# ---------------------------------------------------------

def find_by_name(cursor, table, name, extra_condition="", extra_params=()):
    # Table names and conditions are internal constants, not user input.
    cursor.execute(
        f"""
        SELECT id, is_active
        FROM {table}
        WHERE LOWER(name) = LOWER(%s)
        {extra_condition}
        LIMIT 1
        """,
        (name, *extra_params),
    )

    return cursor.fetchone()


def upsert_named_record(cursor, table, name, is_active):
    existing = find_by_name(cursor, table, name)

    if existing:
        cursor.execute(
            f"""
            UPDATE {table}
            SET
                is_active = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            RETURNING id
            """,
            (is_active, existing["id"]),
        )

        return cursor.fetchone()["id"], "updated"

    cursor.execute(
        f"""
        INSERT INTO {table} (name, is_active)
        VALUES (%s, %s)
        RETURNING id
        """,
        (name, is_active),
    )

    return cursor.fetchone()["id"], "created"


def upsert_service(cursor, group_id, record):
    cursor.execute(
        """
        SELECT id
        FROM services
        WHERE service_group_id = %s
          AND LOWER(name) = LOWER(%s)
        LIMIT 1
        """,
        (group_id, record["name"]),
    )

    existing = cursor.fetchone()

    if existing:
        cursor.execute(
            """
            UPDATE services
            SET
                name = %s,
                available_for_op = %s,
                available_for_ip = %s,
                is_active = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            RETURNING id
            """,
            (
                record["name"],
                record["available_for_op"],
                record["available_for_ip"],
                record["is_active"],
                existing["id"],
            ),
        )

        return cursor.fetchone()["id"], "updated"

    cursor.execute(
        """
        INSERT INTO services (
            service_group_id,
            name,
            available_for_op,
            available_for_ip,
            is_active
        )
        VALUES (%s, %s, %s, %s, %s)
        RETURNING id
        """,
        (
            group_id,
            record["name"],
            record["available_for_op"],
            record["available_for_ip"],
            record["is_active"],
        ),
    )

    return cursor.fetchone()["id"], "created"


def upsert_price(
    cursor,
    service_id,
    bill_type,
    room_type_id,
    doctor_id,
    amount,
    is_active,
):
    cursor.execute(
        """
        SELECT id
        FROM service_prices
        WHERE service_id = %s
          AND bill_type = %s
          AND room_type_id IS NOT DISTINCT FROM %s
          AND doctor_id IS NOT DISTINCT FROM %s
        LIMIT 1
        """,
        (service_id, bill_type, room_type_id, doctor_id),
    )

    existing = cursor.fetchone()

    if existing:
        cursor.execute(
            """
            UPDATE service_prices
            SET
                amount = %s,
                is_active = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s
            """,
            (amount, is_active, existing["id"]),
        )

        return "updated"

    cursor.execute(
        """
        INSERT INTO service_prices (
            service_id,
            bill_type,
            room_type_id,
            doctor_id,
            amount,
            is_active
        )
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (
            service_id,
            bill_type,
            room_type_id,
            doctor_id,
            amount,
            is_active,
        ),
    )

    return "created"


# ---------------------------------------------------------
# IMPORT WORKBOOK IN ONE TRANSACTION
# ---------------------------------------------------------

def import_workbook(data):
    connection = get_db_connection()

    counts = {
        "service_groups_created": 0,
        "service_groups_updated": 0,
        "services_created": 0,
        "services_updated": 0,
        "room_types_created": 0,
        "room_types_updated": 0,
        "doctors_created": 0,
        "doctors_updated": 0,
        "prices_created": 0,
        "prices_updated": 0,
    }

    try:
        with connection:
            with connection.cursor(
                cursor_factory=RealDictCursor
            ) as cursor:

                group_ids = {}
                room_ids = {}
                doctor_ids = {}
                service_ids = {}

                # 1. Service groups
                for row in data["ServiceGroups"]:
                    name = clean_text(row["Service Group Name"])
                    active = parse_yes_no(row["Is Active"], default=True)

                    record_id, action = upsert_named_record(
                        cursor, "service_groups", name, active
                    )

                    group_ids[normalized(name)] = record_id
                    counts[f"service_groups_{action}"] += 1

                # 2. Room types
                for row in data["RoomTypes"]:
                    name = clean_text(row["Room Type Name"])
                    active = parse_yes_no(row["Is Active"], default=True)

                    record_id, action = upsert_named_record(
                        cursor, "room_types", name, active
                    )

                    room_ids[normalized(name)] = record_id
                    counts[f"room_types_{action}"] += 1

                # 3. Doctors
                for row in data["Doctors"]:
                    name = clean_text(row["Doctor Name"])
                    active = parse_yes_no(row["Is Active"], default=True)

                    record_id, action = upsert_named_record(
                        cursor, "doctors", name, active
                    )

                    doctor_ids[normalized(name)] = record_id
                    counts[f"doctors_{action}"] += 1

                # 4. Services
                for row in data["Services"]:
                    group_name = clean_text(row["Service Group Name"])
                    name = clean_text(row["Service Name"])

                    group_id = group_ids.get(normalized(group_name))

                    if group_id is None:
                        existing_group = find_by_name(
                            cursor, "service_groups", group_name
                        )

                        if not existing_group:
                            raise ValueError(
                                f"Service group not found: {group_name}"
                            )

                        group_id = existing_group["id"]

                    record = {
                        "name": name,
                        "available_for_op": parse_yes_no(
                            row["Available for OP"]
                        ),
                        "available_for_ip": parse_yes_no(
                            row["Available for IP"]
                        ),
                        "is_active": parse_yes_no(
                            row["Is Active"], default=True
                        ),
                    }

                    service_id, action = upsert_service(
                        cursor, group_id, record
                    )

                    service_ids[
                        (normalized(group_name), normalized(name))
                    ] = service_id

                    counts[f"services_{action}"] += 1

                # 5. Service prices
                for row in data["ServicePrices"]:
                    group_name = clean_text(row["Service Group Name"])
                    service_name = clean_text(row["Service Name"])
                    bill_type = clean_text(row["Bill Type"]).upper()
                    room_name = clean_text(row["Room Type"])
                    doctor_name = clean_text(row["Doctor Name"])
                    amount = parse_amount(row["Amount"])
                    active = parse_yes_no(
                        row["Is Active"], default=True
                    )

                    service_key = (
                        normalized(group_name),
                        normalized(service_name),
                    )

                    service_id = service_ids.get(service_key)

                    if service_id is None:
                        raise ValueError(
                            "Service could not be resolved: "
                            f"{group_name} / {service_name}"
                        )

                    room_type_id = None
                    doctor_id = None

                    if room_name:
                        room_type_id = room_ids.get(normalized(room_name))

                        if room_type_id is None:
                            existing_room = find_by_name(
                                cursor, "room_types", room_name
                            )

                            if not existing_room:
                                raise ValueError(
                                    f"Room type not found: {room_name}"
                                )

                            room_type_id = existing_room["id"]

                    if doctor_name:
                        doctor_id = doctor_ids.get(normalized(doctor_name))

                        if doctor_id is None:
                            existing_doctor = find_by_name(
                                cursor, "doctors", doctor_name
                            )

                            if not existing_doctor:
                                raise ValueError(
                                    f"Doctor not found: {doctor_name}"
                                )

                            doctor_id = existing_doctor["id"]

                    action = upsert_price(
                        cursor=cursor,
                        service_id=service_id,
                        bill_type=bill_type,
                        room_type_id=room_type_id,
                        doctor_id=doctor_id,
                        amount=amount,
                        is_active=active,
                    )

                    counts[f"prices_{action}"] += 1

        return counts

    except (psycopg2.Error, ValueError):
        logger.exception("Pricing import failed. Transaction rolled back.")

        raise HTTPException(
            status_code=400,
            detail=(
                "Import failed. No changes were committed. "
                "Check the workbook and database schema, then retry."
            ),
        )

    finally:
        connection.close()


# ---------------------------------------------------------
# 1. DOWNLOAD TEMPLATE
# ---------------------------------------------------------

@router.get("/template")
def download_template():
    """
    Download the standard OP/IP pricing import template.
    """

    if not TEMPLATE_PATH.exists():
        raise HTTPException(
            status_code=404,
            detail=(
                "Pricing template not found. Place the file at "
                "Templates/OP_IP_Pricing_Import_Template.xlsx."
            ),
        )

    return FileResponse(
        path=str(TEMPLATE_PATH),
        filename="OP_IP_Pricing_Import_Template.xlsx",
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


# ---------------------------------------------------------
# 2. PREVIEW IMPORT
# ---------------------------------------------------------

@router.post("/preview")
async def preview_import(
    file: UploadFile = File(...),
):
    """
    Validate the workbook without writing to the database.
    """

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are supported.",
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    data, errors = read_workbook(file_bytes)
    errors.extend(validate_workbook(data))

    if errors:
        return {
            "valid": False,
            "message": "Please correct the errors before importing.",
            "errors": errors,
            "error_count": len(errors),
        }

    total_rows = {
        sheet_name: len(records)
        for sheet_name, records in data.items()
    }

    return {
        "valid": True,
        "message": (
            "Workbook validation passed. "
            "Upload the same corrected workbook to the confirm endpoint."
        ),
        "row_counts": total_rows,
        "total_rows": sum(total_rows.values()),
        "errors": [],
    }


# ---------------------------------------------------------
# 3. CONFIRM IMPORT
# ---------------------------------------------------------

@router.post("/confirm")
async def confirm_import(
    file: UploadFile = File(...),
):
    """
    Validate again and import all sheets transactionally.

    The same workbook must be uploaded again after preview.
    """

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are supported.",
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    data, errors = read_workbook(file_bytes)
    errors.extend(validate_workbook(data))

    if errors:
        raise HTTPException(
            status_code=422,
            detail={
                "message": (
                    "Import cancelled because the workbook "
                    "contains validation errors."
                ),
                "errors": errors,
            },
        )

    counts = import_workbook(data)

    return {
        "success": True,
        "message": "Pricing workbook imported successfully.",
        "results": counts,
    }
